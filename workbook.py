import datetime
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from mapping import DATE_LENGTH, DONE_SUFFIX, STATUSES, Record, normalize_status, normalize_who

TASKS_SHEET = "Tasks"
PROJECTS_SHEET = "Projects"
HEADERS = ("Task", "Project", "Start date", "Due date", "Status", "Who", "Follow-up on", "Notes", "Repeat", "Priority", "Link", "ID")
PRIORITIES = ("P1", "P2", "P3", "P4")
LOWEST_PRIORITY = 4
FIRST_DATA_ROW = 2
VALIDATION_LAST_ROW = 5000
DATE_FORMAT = "yyyy-mm-dd"
TASK_URL = "https://app.todoist.com/app/task/{}"
LINK_TEXT = "Open"
BACKUP_STAMP = "%Y%m%d-%H%M%S-%f"
TEMP_SUFFIX = ".tmp"
FORMULA_PREFIX = "="
TEXT_TYPE = "s"
DROPDOWN_HEADERS = ("Project", "Status", "Priority")


@dataclass
class Row:
	id: str | None
	record: Record
	start_date: Any = None
	extra: dict[int, Any] = field(default_factory=dict)


class WorkbookError(Exception):
	pass


def is_open(path):
	path = Path(path)
	lock_names = ("~$" + path.name, "~$" + path.name[2:], f".~lock.{path.name}#")
	if any((path.parent / name).exists() for name in lock_names):
		return True
	if not path.exists():
		return False
	try:
		with open(path, "r+b"):
			return False
	except PermissionError:
		return True


def load_rows(path, owner):
	sheet = load_workbook(path)[TASKS_SHEET]
	columns = _columns(sheet)
	known = set(columns.values())
	rows = []
	for cells in sheet.iter_rows(min_row=FIRST_DATA_ROW):
		values = [cell.value for cell in cells]
		if all(value in (None, "") for value in values):
			continue
		cell = {header: values[index] if index < len(values) else None for header, index in columns.items()}
		_check_link(cells, columns, _text(cell["ID"]))
		record = Record(
			content=_text(cell["Task"]),
			project=_text(cell["Project"]),
			due=_to_iso(cell["Due date"]),
			status=normalize_status(cell["Status"]),
			who=normalize_who(_text(cell["Who"]), owner),
			follow_up=_to_iso(cell["Follow-up on"]),
			notes=_text(cell["Notes"]),
			repeat=_text(cell["Repeat"]) or None,
			priority=_priority(cell["Priority"]),
		)
		extra = {index: value for index, value in enumerate(values) if index not in known and value is not None}
		rows.append(Row(_text(cell["ID"]) or None, record, cell["Start date"], extra))
	return rows


def save_rows(path, rows, project_names, owner):
	path = Path(path)
	if path.exists():
		book = load_workbook(path)
		sheet = book[TASKS_SHEET]
	else:
		book = Workbook()
		sheet = book.active
		sheet.title = TASKS_SHEET
		sheet.append(HEADERS)
	columns = _columns(sheet)
	_clear_data(sheet)
	for row_number, row in enumerate(rows, start=FIRST_DATA_ROW):
		_write_row(sheet, row_number, columns, row, owner)
	_write_projects(book, project_names)
	_add_dropdowns(sheet, columns, len(project_names))
	sheet.column_dimensions[_letter(columns, "ID")].hidden = True
	_save_atomic(book, path)


def backup(path, backup_dir, keep):
	path = Path(path)
	if not path.exists():
		return
	backup_dir = Path(backup_dir)
	backup_dir.mkdir(exist_ok=True)
	stamp = datetime.datetime.now().strftime(BACKUP_STAMP)
	shutil.copy2(path, backup_dir / f"{path.stem}-{stamp}{path.suffix}")
	for old in sorted(backup_dir.glob(f"{path.stem}-*{path.suffix}"))[:-keep]:
		old.unlink()


def _columns(sheet):
	header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
	columns = {name: index for index, name in enumerate(header_row) if name in HEADERS}
	missing = [header for header in HEADERS if header not in columns]
	if missing:
		raise WorkbookError(f"Missing columns in {TASKS_SHEET}: {', '.join(missing)}")
	return columns


def _check_link(cells, columns, task_id):
	index = columns["Link"]
	link = cells[index].hyperlink if index < len(cells) else None
	if not task_id or link is None:
		return
	if link.target != TASK_URL.format(task_id.removesuffix(DONE_SUFFIX)):
		raise WorkbookError(
			f"Row {cells[0].row}: the Link does not match the hidden ID. "
			"The rows were probably sorted or filtered without the hidden ID column. "
			"Undo that, or restore a copy from the backups folder."
		)


def _clear_data(sheet):
	for cells in sheet.iter_rows(min_row=FIRST_DATA_ROW):
		for cell in cells:
			cell.hyperlink = None
			cell.value = None


def _save_atomic(book, path):
	temp = path.with_name(path.name + TEMP_SUFFIX)
	try:
		book.save(temp)
		temp.replace(path)
	finally:
		temp.unlink(missing_ok=True)


def _letter(columns, header):
	return get_column_letter(columns[header] + 1)


def _write_row(sheet, row_number, columns, row, owner):
	record = row.record
	values = {
		"Task": record.content,
		"Project": record.project,
		"Start date": row.start_date,
		"Due date": _to_date(record.due),
		"Status": record.status,
		"Who": record.who or owner,
		"Follow-up on": _to_date(record.follow_up),
		"Notes": record.notes,
		"Repeat": record.repeat,
		"Priority": f"P{record.priority}",
		"Link": None,
		"ID": row.id,
	}
	for header, value in values.items():
		cell = _set_cell(sheet, row_number, columns[header], value)
		if isinstance(value, str) and value.startswith(FORMULA_PREFIX):
			cell.data_type = TEXT_TYPE
		if isinstance(value, datetime.date):
			cell.number_format = DATE_FORMAT
	for index, value in row.extra.items():
		_set_cell(sheet, row_number, index, value)
	if row.id:
		link = sheet.cell(row=row_number, column=columns["Link"] + 1, value=LINK_TEXT)
		link.hyperlink = TASK_URL.format(row.id.removesuffix(DONE_SUFFIX))
		link.style = "Hyperlink"


def _set_cell(sheet, row_number, index, value):
	if isinstance(value, str):
		value = ILLEGAL_CHARACTERS_RE.sub("", value)
	return sheet.cell(row=row_number, column=index + 1, value=value)


def _write_projects(book, project_names):
	if PROJECTS_SHEET in book.sheetnames:
		del book[PROJECTS_SHEET]
	sheet = book.create_sheet(PROJECTS_SHEET)
	for name in project_names:
		sheet.append([name])
	sheet.sheet_state = "hidden"


def _add_dropdowns(sheet, columns, project_count):
	owned = {columns[header] + 1 for header in DROPDOWN_HEADERS}
	sheet.data_validations.dataValidation = [
		validation for validation in sheet.data_validations.dataValidation
		if not _is_owned(validation, owned)
	]
	lists = {
		"Project": f"{PROJECTS_SHEET}!$A$1:$A${max(project_count, 1)}",
		"Status": '"' + ",".join(STATUSES) + '"',
		"Priority": '"' + ",".join(PRIORITIES) + '"',
	}
	for header, formula in lists.items():
		validation = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True)
		letter = _letter(columns, header)
		validation.add(f"{letter}{FIRST_DATA_ROW}:{letter}{VALIDATION_LAST_ROW}")
		sheet.add_data_validation(validation)


def _is_owned(validation, owned):
	return all(cells.min_col == cells.max_col and cells.min_col in owned for cells in validation.sqref.ranges)


def _text(value):
	return "" if value is None else str(value).strip()


def _to_iso(value):
	if isinstance(value, datetime.datetime):
		return value.date().isoformat()
	if isinstance(value, datetime.date):
		return value.isoformat()
	return _text(value)[:DATE_LENGTH] or None


def _to_date(iso):
	return datetime.date.fromisoformat(iso) if iso else None


def _priority(value):
	text = _text(value).upper()
	return PRIORITIES.index(text) + 1 if text in PRIORITIES else LOWEST_PRIORITY
