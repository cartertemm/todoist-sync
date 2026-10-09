import datetime

import pytest
from openpyxl import Workbook, load_workbook

from mapping import Record
from workbook import Row, WorkbookError, backup, is_open, load_rows, save_rows

OWNER = "Carter"
PROJECTS = ["Inbox", "Work", "Home"]


def sample_rows():
	return [
		Row(
			"t1",
			Record(
				content="Write report",
				project="Work",
				due="2026-10-10",
				status="Blocked",
				who="Jane",
				follow_up="2026-10-12",
				notes="Ask about budget",
				repeat=None,
				priority=1,
			),
			start_date=datetime.datetime(2026, 10, 1),
		),
		Row(
			"t2",
			Record(
				content="Water plants",
				project="Home",
				due=None,
				status="Todo",
				who=None,
				follow_up=None,
				notes="",
				repeat="every week",
				priority=4,
			),
		),
	]


@pytest.fixture
def path(tmp_path):
	return tmp_path / "Tasks.xlsx"


def test_save_then_load_round_trips(path):
	save_rows(path, sample_rows(), PROJECTS, OWNER)
	assert load_rows(path, OWNER) == sample_rows()


def test_blank_who_is_written_as_owner(path):
	save_rows(path, sample_rows(), PROJECTS, OWNER)
	assert load_workbook(path)["Tasks"]["F3"].value == OWNER


def test_link_points_at_task_and_history_rows_link_to_parent(path):
	rows = sample_rows() + [Row("t2:done", sample_rows()[1].record)]
	save_rows(path, rows, PROJECTS, OWNER)
	sheet = load_workbook(path)["Tasks"]
	assert sheet["K2"].hyperlink.target == "https://app.todoist.com/app/task/t1"
	assert sheet["K4"].hyperlink.target == "https://app.todoist.com/app/task/t2"


def test_id_column_and_projects_sheet_are_hidden(path):
	save_rows(path, sample_rows(), PROJECTS, OWNER)
	book = load_workbook(path)
	assert book["Tasks"].column_dimensions["L"].hidden
	assert book["Projects"].sheet_state == "hidden"
	assert [cell.value for cell in book["Projects"]["A"]] == PROJECTS


def test_dropdowns_are_added_once(path):
	save_rows(path, sample_rows(), PROJECTS, OWNER)
	save_rows(path, sample_rows(), PROJECTS, OWNER)
	validations = load_workbook(path)["Tasks"].data_validations.dataValidation
	assert sorted(validation.formula1 for validation in validations) == sorted([
		"Projects!$A$1:$A$3",
		'"Todo,Complete,Blocked"',
		'"P1,P2,P3,P4"',
	])


def test_saving_fewer_rows_removes_old_rows(path):
	save_rows(path, sample_rows(), PROJECTS, OWNER)
	save_rows(path, sample_rows()[:1], PROJECTS, OWNER)
	assert load_rows(path, OWNER) == sample_rows()[:1]


def test_empty_rows_are_skipped(path):
	save_rows(path, sample_rows(), PROJECTS, OWNER)
	book = load_workbook(path)
	book["Tasks"].append([None] * 12)
	book["Tasks"].append(["Typed in Excel", "Work"])
	book.save(path)
	rows = load_rows(path, OWNER)
	assert len(rows) == 3
	assert rows[2].id is None
	assert rows[2].record.content == "Typed in Excel"
	assert rows[2].record.priority == 4


def test_missing_column_raises(path):
	book = Workbook()
	book.active.title = "Tasks"
	book.active.append(["Task"])
	book.save(path)
	with pytest.raises(WorkbookError, match="Project"):
		load_rows(path, OWNER)


def test_is_open_detects_lock_files(path):
	save_rows(path, sample_rows(), PROJECTS, OWNER)
	assert not is_open(path)
	lock = path.parent / "~$Tasks.xlsx"
	lock.write_text("")
	assert is_open(path)
	lock.unlink()
	(path.parent / ".~lock.Tasks.xlsx#").write_text("")
	assert is_open(path)


def test_backup_keeps_newest_copies(path, tmp_path):
	save_rows(path, sample_rows(), PROJECTS, OWNER)
	backup_dir = tmp_path / "backups"
	for _ in range(12):
		backup(path, backup_dir, keep=10)
	assert len(list(backup_dir.iterdir())) == 10


def test_backup_without_workbook_does_nothing(path, tmp_path):
	backup(path, tmp_path / "backups", keep=10)
	assert not (tmp_path / "backups").exists()
