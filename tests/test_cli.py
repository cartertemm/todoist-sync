import datetime
from dataclasses import replace

import pytest
from openpyxl import load_workbook

import sync
from fakes import FakeTodoist
from workbook import load_rows, save_rows

PROJECTS = ["Inbox", "Work", "Home"]
NOW = datetime.datetime(2026, 10, 9, 12, 0, tzinfo=datetime.timezone.utc)


@pytest.fixture
def setup(tmp_path, monkeypatch):
	workbook_path = tmp_path / "Tasks.xlsx"
	config_path = tmp_path / "config.toml"
	config_path.write_text(
		f'todoist_token = "from-file"\nworkbook = "{workbook_path.as_posix()}"\nowner = "Carter"\n',
		encoding="utf-8",
	)
	fake = FakeTodoist()
	monkeypatch.setattr(sync, "STATE_PATH", tmp_path / "state.json")
	monkeypatch.setattr(sync, "LOG_PATH", tmp_path / "sync.log")
	monkeypatch.setattr(sync, "BACKUP_DIR", tmp_path / "backups")
	monkeypatch.setattr(sync, "TodoistClient", lambda token: fake)
	monkeypatch.delenv(sync.TOKEN_ENV, raising=False)
	return fake, workbook_path, config_path, tmp_path


def test_env_token_overrides_file(setup, monkeypatch):
	_, _, config_path, _ = setup
	assert sync.load_config(config_path)["token"] == "from-file"
	monkeypatch.setenv(sync.TOKEN_ENV, "from-env")
	config = sync.load_config(config_path)
	assert config["token"] == "from-env"
	assert config["max_delete_percent"] == sync.DEFAULT_MAX_DELETE_PERCENT


def test_state_round_trip(tmp_path):
	path = tmp_path / "state.json"
	assert sync.load_state(path) == {}
	sync.save_state(path, {"last_sync": "x", "tasks": {}})
	assert sync.load_state(path) == {"last_sync": "x", "tasks": {}}


def test_first_run_creates_workbook_and_ignores_old_state(setup):
	fake, workbook_path, config_path, tmp_path = setup
	fake.add("Write report")
	sync.save_state(sync.STATE_PATH, {"last_sync": NOW.isoformat(), "tasks": {"gone": {}}})
	assert sync.main(["--config", str(config_path)]) == 0
	assert [row.record.content for row in load_rows(workbook_path, "Carter")] == ["Write report"]
	assert fake.deleted == []
	assert "gone" not in sync.load_state(sync.STATE_PATH)["tasks"]


def test_open_workbook_is_skipped(setup):
	fake, workbook_path, config_path, tmp_path = setup
	workbook_path.write_bytes(b"")
	(tmp_path / "~$Tasks.xlsx").write_text("")
	assert sync.main(["--config", str(config_path)]) == 0
	assert "workbook is open" in sync.LOG_PATH.read_text(encoding="utf-8")
	assert fake.calls == []


def test_delete_guard_prints_tasks_and_changes_nothing(setup, capsys):
	fake, workbook_path, config_path, tmp_path = setup
	for number in range(4):
		fake.add(f"Task {number}")
	sync.main(["--config", str(config_path)])
	rows = load_rows(workbook_path, "Carter")
	sync.workbook.save_rows(workbook_path, rows[2:], ["Inbox", "Work", "Home"], "Carter")
	assert sync.main(["--config", str(config_path)]) == 1
	output = capsys.readouterr().out
	assert "Task 0" in output and "--allow-deletes" in output
	assert fake.deleted == []
	assert sync.main(["--config", str(config_path), "--allow-deletes"]) == 0
	assert len(fake.deleted) == 2


def test_each_run_writes_one_log_line(setup):
	fake, _, config_path, _ = setup
	fake.add("Write report")
	sync.main(["--config", str(config_path)])
	sync.main(["--config", str(config_path)])
	lines = sync.LOG_PATH.read_text(encoding="utf-8").splitlines()
	assert len(lines) == 2
	assert "added row: Write report" in lines[0]
	assert lines[1].endswith("no changes")


def fail_save(*args):
	raise PermissionError("the workbook is open")


def run_with_failed_save(monkeypatch, config_path):
	before = sync.STATE_PATH.read_text(encoding="utf-8")
	with monkeypatch.context() as patch:
		patch.setattr(sync.workbook, "save_rows", fail_save)
		assert sync.main(["--config", str(config_path)]) == 1
	assert sync.STATE_PATH.read_text(encoding="utf-8") == before
	assert "workbook not saved" in sync.LOG_PATH.read_text(encoding="utf-8")


def test_failed_save_then_clean_run_creates_task_once(setup, monkeypatch):
	fake, workbook_path, config_path, _ = setup
	sync.main(["--config", str(config_path)])
	book = load_workbook(workbook_path)
	book["Tasks"].append(["New thing", "Work"])
	book.save(workbook_path)
	run_with_failed_save(monkeypatch, config_path)
	assert sync.main(["--config", str(config_path)]) == 0
	assert len(fake.tasks) == 1
	assert [row.id for row in load_rows(workbook_path, "Carter")] == list(fake.tasks)


def test_failed_save_then_clean_run_closes_recurring_task_once(setup, monkeypatch):
	fake, workbook_path, config_path, _ = setup
	task_id = fake.add("Water plants", due={"date": "2026-10-09", "string": "every week", "is_recurring": True})
	sync.main(["--config", str(config_path)])
	rows = load_rows(workbook_path, "Carter")
	rows[0] = replace(rows[0], record=replace(rows[0].record, status="Complete"))
	save_rows(workbook_path, rows, PROJECTS, "Carter")
	run_with_failed_save(monkeypatch, config_path)
	assert sync.main(["--config", str(config_path)]) == 0
	assert [call for call in fake.calls if call[0] == "close_task"] == [("close_task", task_id)]
	assert [row.id for row in load_rows(workbook_path, "Carter")] == [task_id, f"{task_id}:done"]


def test_misaligned_link_stops_run(setup):
	fake, workbook_path, config_path, _ = setup
	fake.add("A")
	fake.add("B")
	sync.main(["--config", str(config_path)])
	book = load_workbook(workbook_path)
	sheet = book["Tasks"]
	sheet["K2"].hyperlink, sheet["K3"].hyperlink = sheet["K3"].hyperlink.target, sheet["K2"].hyperlink.target
	book.save(workbook_path)
	fake.calls.clear()
	assert sync.main(["--config", str(config_path)]) == 1
	assert fake.calls == []
	assert "Row 2" in sync.LOG_PATH.read_text(encoding="utf-8")


def test_unexpected_error_is_logged_with_traceback(setup, monkeypatch):
	_, _, config_path, _ = setup

	def explode(*args):
		raise ValueError("boom")

	monkeypatch.setattr(sync, "run_sync", explode)
	assert sync.main(["--config", str(config_path)]) == 1
	lines = sync.LOG_PATH.read_text(encoding="utf-8").splitlines()
	assert lines[0].endswith("error: ValueError: boom")
	assert lines[1].startswith("\tTraceback")
